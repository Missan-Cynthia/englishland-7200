"""Import the supplied Unit 1–30 document; never merge legacy PDF content."""
import json
import re
import sys
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]

def extract(source):
    document = Document(source)
    assert len(document.tables) == 62, 'Unexpected document structure'
    phrases, units = [], []
    for unit in range(1, 31):
        passage = document.tables[unit * 2].cell(0, 0).text.strip()
        table = document.tables[unit * 2 + 1]
        assert [c.text for c in table.rows[0].cells] == ['編號', '片語', '中文']
        sentences = [s.strip() for s in re.split(r'(?<=[.!?])\s+|\n+', passage) if s.strip()]
        ids = []
        for number, row in enumerate(table.rows[1:], 1):
            ordinal, phrase, meaning = [c.text.strip() for c in row.cells]
            assert ordinal == str(number) and phrase and meaning
            identifier = f'docx-u{unit:02d}-{number:02d}'
            forms = [re.sub(r'^to\s+', '', f.strip()) for f in phrase.split('/')]
            examples = [s for s in sentences if any(re.search(r'\b' + re.escape(f) + r'\b', s, re.I) for f in forms)]
            phrases.append(dict(id=identifier, lesson=unit, phrase=phrase, meaningZh=meaning,
                                meaningEn='', examples=examples[:3]))
            ids.append(identifier)
        units.append(dict(unit=unit, passage=passage, phraseIds=ids))
    return phrases, units

if __name__ == '__main__':
    phrases, units = extract(sys.argv[1])
    for name, data in [('phrases.json', phrases), ('phrase-units.json', units)]:
        (ROOT / 'data' / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Imported {len(phrases)} phrases and {len(units)} units')
